import asyncio
import csv
import json
import os
import tempfile
import unittest
from pathlib import Path

from docx import Document
from mcp import Client, StdioServerParameters

from pentwin.cli import CorpusError, prepare, seal
from pentwin.server import audit_draft, build_server, load_corpus
from pentwin.style_engine import aggregate


class MCPTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        papers = []
        for i in range(4):
            doc = Document()
            doc.add_heading("Introduction", 1)
            doc.add_paragraph(f"A synthetic pathway in set {i} may alter expression. This preliminary finding requires an independent assay to assess the response.")
            doc.add_paragraph(f"In set {i}, the control and treatment samples were assayed separately. The observations are limited to these artificial measurements.")
            source = self.root / f"sample{i}.docx"
            doc.save(source)
            papers.append({"paper_id": f"p{i}", "family_id": f"f{i}", "file": source.name,
                           "authorship": "primary_confirmed", "status": "published", "model_excerpt_permission": "yes"})
        self.manifest = self.root / "manifest.json"
        self.manifest.write_text(json.dumps({"papers": papers}))
        self.run = self.root / "run"
        prepare(self.manifest, self.run, "synthetic")
        with (self.run / "review.csv").open(newline="") as fh:
            rows = list(csv.DictReader(fh))
        for row in rows:
            row.update(include="yes", source_verified="yes", section="introduction", move="background")
        with (self.run / "review.csv").open("w", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=rows[0].keys())
            writer.writeheader()
            writer.writerows(rows)
        seal(self.run)

    def tearDown(self):
        self.temp.cleanup()

    def test_permission_no_still_allows_aggregate_but_no_source_text_leaves(self):
        rows = load_corpus(self.run)
        held = [r for r in rows if r["split"] == "holdout"]
        denied = next(r for r in rows if r["split"] == "development")
        unpublished = next(r for r in rows if r["split"] == "development" and r["family_id"] != denied["family_id"])
        for row in rows:
            if row["family_id"] == denied["family_id"]:
                row["model_excerpt_permission"] = "no"
            if row["family_id"] == unpublished["family_id"]:
                row["status"] = "unpublished"
        response = json.dumps(aggregate(rows))
        self.assertNotIn(held[0]["text"], response)
        self.assertNotIn(denied["text"], response)
        self.assertNotIn(unpublished["text"], response)
        self.assertEqual(aggregate(rows)["status"], "provisional")

    def test_tampered_corpus_and_unsealed_run_fail_closed(self):
        (self.run / "corpus.jsonl").write_text("{}\n")
        with self.assertRaisesRegex(CorpusError, "changed"):
            load_corpus(self.run)
        with self.assertRaises(CorpusError):
            build_server(self.root)

    def test_fact_alarm_flags_changed_number_without_claiming_approval(self):
        result = audit_draft("The response was 12% in treated samples.", "The response was 21% in treated samples.")
        self.assertEqual(result["status"], "needs_human_review")
        self.assertIn("21%", result["new_numeric_tokens"])
        self.assertIn("12%", result["omitted_numeric_tokens"])
        self.assertEqual(audit_draft("Treatment did not improve survival.", "Treatment improved survival.")["status"], "needs_human_review")

    def test_stdio_sdk_handshake_discovery_and_tool_call(self):
        async def exercise():
            import sys
            server = StdioServerParameters(command=sys.executable, args=["-m", "pentwin.cli", "serve", str(self.run)],
                                           env={**os.environ, "PYTHONPATH": str(Path(__file__).resolve().parents[1] / "src")})
            async with Client(server) as client:
                tools = await client.list_tools()
                self.assertEqual({t.name for t in tools.tools}, {"corpus_status", "style_profile", "review_draft", "rewrite_plan", "assess_revision"})
                status = await client.call_tool("corpus_status", {})
                self.assertFalse(status.is_error)
                self.assertTrue(json.loads(status.content[0].text)["sealed"])
                profile = await client.call_tool("style_profile", {})
                self.assertFalse(profile.is_error)
                response = json.loads(profile.content[0].text)
                self.assertEqual(response["status"], "provisional")
                self.assertFalse(any(k.startswith("function:") for k in response["target"]))
                held = [r["text"] for r in load_corpus(self.run) if r["split"] == "holdout"]
                self.assertTrue(all(t not in json.dumps(response) for t in held))
                self.assertTrue(all(r["text"] not in json.dumps(response) for r in load_corpus(self.run)))
                plan = await client.call_tool("rewrite_plan", {"draft": "The synthetic response was 12.5 mg/mL in the assay.",
                                                               "section": "introduction", "protected_spans": ["12.5 mg/mL"]})
                self.assertFalse(plan.is_error)
                plan_response = json.loads(plan.content[0].text)
                self.assertEqual(plan_response["status"], "rewrite_guidance")
                self.assertFalse(any(k.startswith("function:") for k in plan_response["target"]))
                self.assertEqual(plan_response["protected_spans"], ["12.5 mg/mL"])
                checked = await client.call_tool("assess_revision", {
                    "original": "The synthetic response was 12.5 mg/mL in the assay.",
                    "revised": "The synthetic response was 15 mg/mL in the assay.",
                    "section": "introduction", "protected_spans": ["12.5 mg/mL"]})
                self.assertFalse(checked.is_error)
                self.assertEqual(json.loads(checked.content[0].text)["protected_span_changes"][0]["span"], "12.5 mg/mL")
        asyncio.run(exercise())


if __name__ == "__main__":
    unittest.main()
