import csv
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas
from pentwin.cli import CorpusError, _pdf_candidates, prepare, seal


class StageOneTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        papers = []
        for i in range(3):
            doc = Document()
            doc.add_heading("Introduction", 1)
            doc.add_paragraph(f"The biological pathway in study {i} may affect the observed response in this experiment.")
            doc.add_heading("References", 1)
            doc.add_paragraph("This cited paragraph must not enter the corpus.")
            path = self.root / f"p{i}.docx"
            doc.save(path)
            papers.append({"paper_id": f"p{i}", "family_id": f"f{i}", "file": path.name,
                           "authorship": "primary_confirmed", "status": "published",
                           "model_excerpt_permission": "no"})
        self.manifest = self.root / "manifest.json"
        self.manifest.write_text(json.dumps({"papers": papers}), encoding="utf-8")
        self.out = self.root / "run"

    def tearDown(self):
        self.tmp.cleanup()

    def _rows(self):
        with (self.out / "review.csv").open(encoding="utf-8", newline="") as fh:
            return list(csv.DictReader(fh))

    def _write_rows(self, rows):
        with (self.out / "review.csv").open("w", encoding="utf-8", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=rows[0].keys())
            writer.writeheader()
            writer.writerows(rows)

    def test_seal_requires_explicit_review_and_preserves_holdout(self):
        result = prepare(self.manifest, self.out, "fixed")
        self.assertEqual(result["candidates"], 3)
        self.assertEqual(len({r["family_id"] for r in self._rows() if r["split"] == "holdout"}), 1)
        with self.assertRaises(CorpusError):
            seal(self.out)
        rows = self._rows()
        for row in rows:
            row["include"] = "yes"
            row["source_verified"] = "yes"
            row["section"] = "introduction"
            row["move"] = "background"
        self._write_rows(rows)
        self.assertEqual(seal(self.out)["accepted"], 3)
        self.assertEqual(len((self.out / "corpus.jsonl").read_text().splitlines()), 3)
        self.assertNotIn("cited paragraph", (self.out / "corpus.jsonl").read_text())
        with self.assertRaises(CorpusError):
            seal(self.out)

    def test_rejects_modified_extracted_text_and_missing_row(self):
        prepare(self.manifest, self.out, "fixed")
        rows = self._rows()
        for row in rows:
            row.update(include="yes", source_verified="yes", section="introduction", move="background")
        rows[0]["text"] = "Tampered text that was not in the source paper."
        self._write_rows(rows)
        with self.assertRaisesRegex(CorpusError, "changed"):
            seal(self.out)
        self._write_rows(rows[1:])
        with self.assertRaises(CorpusError):
            seal(self.out)

    def test_same_family_cannot_create_fake_independent_holdout(self):
        papers = json.loads(self.manifest.read_text())["papers"]
        papers[1]["family_id"] = papers[0]["family_id"]
        self.manifest.write_text(json.dumps({"papers": papers}))
        with self.assertRaisesRegex(CorpusError, "independent"):
            prepare(self.manifest, self.out, "fixed")

    def test_duplicate_bytes_across_families_are_rejected(self):
        (self.root / "p1.docx").write_bytes((self.root / "p0.docx").read_bytes())
        with self.assertRaisesRegex(CorpusError, "Identical source"):
            prepare(self.manifest, self.out, "fixed")

    def test_private_output_inside_git_tree_is_rejected(self):
        (self.root / ".git").mkdir()
        with self.assertRaisesRegex(CorpusError, "outside every Git"):
            prepare(self.manifest, self.out, "fixed")

    def test_changed_source_cannot_be_sealed(self):
        prepare(self.manifest, self.out, "fixed")
        doc = Document(self.root / "p0.docx")
        doc.add_paragraph("A changed paragraph invalidates the reviewed extraction.")
        doc.save(self.root / "p0.docx")
        with self.assertRaisesRegex(CorpusError, "changed"):
            seal(self.out)

    @unittest.skipUnless(shutil.which("pandoc"), "pandoc is required for LaTeX extraction")
    def test_pdf_and_latex_are_candidates_needing_review(self):
        pdf = canvas.Canvas(str(self.root / "p1.pdf"))
        pdf.drawString(50, 750, "A biological response was observed in the sample cohort.")
        pdf.save()
        (self.root / "p2.tex").write_text(
            r"\documentclass{article}\begin{document}\section{Results}The assay showed a biological response across three experimental conditions.\end{document}",
            encoding="utf-8",
        )
        papers = json.loads(self.manifest.read_text())["papers"]
        papers[1]["file"] = "p1.pdf"
        papers[2]["file"] = "p2.tex"
        self.manifest.write_text(json.dumps({"papers": papers}))
        prepare(self.manifest, self.out, "fixed")
        rows = self._rows()
        self.assertEqual({r["paper_id"] for r in rows}, {"p0", "p1", "p2"})
        self.assertTrue(all(not r["source_verified"] for r in rows))
        self.assertIn("pdf_order", next(r["extraction_warning"] for r in rows if r["paper_id"] == "p1"))

    def test_unconfirmed_authorship_cannot_be_accepted(self):
        papers = json.loads(self.manifest.read_text())["papers"]
        papers[0]["authorship"] = "uncertain"
        self.manifest.write_text(json.dumps({"papers": papers}))
        prepare(self.manifest, self.out, "fixed")
        rows = self._rows()
        for row in rows:
            row.update(include="yes", source_verified="yes", section="introduction", move="background")
        self._write_rows(rows)
        with self.assertRaisesRegex(CorpusError, "authorship"):
            seal(self.out)

    def test_pdf_section_continues_across_pages_and_references_stop(self):
        source = self.root / "multipage.pdf"
        pdf = canvas.Canvas(str(source))
        pdf.drawString(50, 750, "Introduction")
        pdf.drawString(50, 730, "The first synthetic paragraph establishes an artificial example for this test.")
        pdf.showPage()
        pdf.drawString(50, 750, "The second synthetic paragraph continues this introduction on the next page.")
        pdf.drawString(50, 720, "References")
        pdf.showPage()
        pdf.drawString(50, 750, "This cited reference text must be excluded on all later pages.")
        pdf.save()
        rows = list(_pdf_candidates(source))
        self.assertEqual(len(rows), 2)
        self.assertEqual([r[1] for r in rows], ["introduction", "introduction"])
        self.assertNotIn("cited reference", " ".join(r[2] for r in rows))


if __name__ == "__main__":
    unittest.main()
